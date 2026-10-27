#!/usr/bin/env python3
"""Recalculate every formula of a workbook with LibreOffice and save the result to a new file.

Usage: python3 recalc.py <input.xlsx> <output.xlsx> [timeout_seconds]
Prints one JSON object:
  {"ok": bool, "formulas": int, "errors": [{"cell": "Sheet!B4", "value": "#REF!"}]}
or, when nothing could be recalculated, {"ok": false, "error": "..."}.

LibreOffice runs headless with a profile of its own, driven through its Python bridge (python3-uno):
it loads the workbook, calculates all of it, and saves it as Excel 2007+.
"""
import json
import os
import shutil
import subprocess
import sys
import tempfile
import time

ERROR_VALUES = {"#REF!", "#NAME?", "#VALUE!", "#DIV/0!", "#N/A", "#NUM!", "#NULL!"}


def fail(message):
    print(json.dumps({"ok": False, "error": message}, ensure_ascii=False))
    sys.exit(1)


def file_url(path):
    import uno

    return uno.systemPathToFileUrl(os.path.abspath(path))


def prop(name, value):
    from com.sun.star.beans import PropertyValue

    p = PropertyValue()
    p.Name = name
    p.Value = value
    return p


def recalculate(source, target, timeout):
    try:
        import uno
        from com.sun.star.connection import NoConnectException
    except ImportError:
        fail("LibreOffice's Python bridge (python3-uno) is not installed in this terminal")
    if not shutil.which("soffice"):
        fail("LibreOffice (soffice) is not installed in this terminal")
    profile = tempfile.mkdtemp(prefix="lo-profile-", dir="/tmp")
    pipe = f"recalc{os.getpid()}"
    office = subprocess.Popen(
        [
            "soffice", "--headless", "--invisible", "--nologo", "--norestore", "--nodefault",
            f"-env:UserInstallation=file://{profile}",
            f"--accept=pipe,name={pipe};urp;StarOffice.ComponentContext",
        ],
        stdout=subprocess.DEVNULL,
        stderr=subprocess.DEVNULL,
    )
    try:
        local = uno.getComponentContext()
        resolver = local.ServiceManager.createInstanceWithContext("com.sun.star.bridge.UnoUrlResolver", local)
        deadline = time.time() + timeout
        while True:
            try:
                ctx = resolver.resolve(f"uno:pipe,name={pipe};urp;StarOffice.ComponentContext")
                break
            except NoConnectException:
                if time.time() > deadline or office.poll() is not None:
                    fail("LibreOffice did not start in time")
                time.sleep(0.3)
        desktop = ctx.ServiceManager.createInstanceWithContext("com.sun.star.frame.Desktop", ctx)
        doc = desktop.loadComponentFromURL(file_url(source), "_blank", 0, (prop("Hidden", True),))
        if doc is None:
            fail(f"LibreOffice could not open {source}")
        try:
            doc.calculateAll()
            doc.storeToURL(file_url(target), (prop("FilterName", "Calc MS Excel 2007 XML"),))
        finally:
            doc.close(True)
    finally:
        office.terminate()
        try:
            office.wait(10)
        except subprocess.TimeoutExpired:
            office.kill()
        shutil.rmtree(profile, ignore_errors=True)


def scan(target):
    from openpyxl import load_workbook

    formulas = 0
    with_formulas = load_workbook(target, data_only=False)
    for sheet in with_formulas.worksheets:
        for row in sheet.iter_rows():
            formulas += sum(1 for c in row if isinstance(c.value, str) and c.value.startswith("="))
    errors = []
    values = load_workbook(target, data_only=True)
    for sheet in values.worksheets:
        for row in sheet.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value in ERROR_VALUES:
                    errors.append({"cell": f"{sheet.title}!{c.coordinate}", "value": c.value})
    return formulas, errors


if __name__ == "__main__":
    if len(sys.argv) not in (3, 4):
        fail("usage: recalc.py <input.xlsx> <output.xlsx> [timeout_seconds]")
    source, target = sys.argv[1], sys.argv[2]
    if not os.path.isfile(source):
        fail(f"{source} does not exist")
    if os.path.abspath(source) == os.path.abspath(target):
        fail("write the recalculated workbook to a new file")
    recalculate(source, target, float(sys.argv[3]) if len(sys.argv) == 4 else 60.0)
    formulas, errors = scan(target)
    print(json.dumps({"ok": not errors, "formulas": formulas, "errors": errors}, ensure_ascii=False))
    sys.exit(0 if not errors else 1)
